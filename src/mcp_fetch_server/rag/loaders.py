"""Turn files into :class:`NormalizedDoc` objects.

Each loader's job is to recover *structure*, not just text. Headings and page
numbers are what let a chunk carry a heading path and a page anchor, and those
are what make a citation checkable. A loader that returns one giant paragraph
technically works and quietly ruins retrieval.

Formats are handled with permissively licensed libraries only. PyMuPDF is
deliberately not used (AGPL), and EPUB is unzipped and parsed with lxml rather
than pulling in ebooklib (also AGPL).
"""

from __future__ import annotations

import logging
import posixpath
import re
import zipfile
from collections.abc import Callable
from pathlib import Path

from lxml import etree
from lxml import html as lxml_html

from mcp_fetch_server.rag.documents import Block, NormalizedDoc

logger = logging.getLogger(__name__)


class LoaderError(Exception):
    """Raised when a file cannot be loaded."""


class UnsupportedFormatError(LoaderError):
    """Raised when no loader is registered for a file extension."""


EXTENSION_DOCTYPES: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".epub": "epub",
    ".html": "html",
    ".htm": "html",
    ".xhtml": "html",
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
    ".text": "text",
    ".rst": "text",
}

SUPPORTED_EXTENSIONS = frozenset(EXTENSION_DOCTYPES)

# Persian, Arabic and Hebrew ranges, used for a cheap language guess.
_RTL_RE = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷏ﹰ-﻿֐-׿]")
_PERSIAN_ONLY_RE = re.compile(r"[پچژگ]")
_CJK_RE = re.compile(r"[぀-ヿ㐀-䶿一-鿿가-힯]")

_BLANK_LINE_RE = re.compile(r"\n\s*\n")
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_FENCE_RE = re.compile(r"^\s*```")
_NUMBERED_HEADING_RE = re.compile(r"^\s*(?:\d+|[IVXLC]+)(?:[.\-]\d+)*[.\)]?\s+\S")


def detect_language(text: str) -> str:
    """Cheap script-based language guess.

    This is not language identification: it separates the script families that
    matter for retrieval (so a Persian document is not filtered as English)
    without adding a model dependency. Refined later during enrichment, where
    the local model sees the actual text.
    """
    sample = text[:4000]
    if not sample.strip():
        return "und"
    letters = sum(1 for char in sample if char.isalpha())
    if not letters:
        return "und"
    if _CJK_RE.search(sample):
        return "zh"
    rtl = len(_RTL_RE.findall(sample))
    if rtl / max(letters, 1) > 0.25:
        return "fa" if _PERSIAN_ONLY_RE.search(sample) else "ar"
    return "en"


def _looks_like_heading(line: str) -> int | None:
    """Guess a heading level for a bare line of text, or None.

    Kept deliberately conservative. A false heading fragments a section and
    costs a little context; treating every short line as a heading would
    shred the document, so the checks require several signals at once.
    """
    stripped = line.strip()
    if not stripped or len(stripped) > 90:
        return None
    if stripped.endswith((".", ",", ";", ":", "،", "؛")):
        return None
    words = stripped.split()
    if len(words) > 14:
        return None

    if _NUMBERED_HEADING_RE.match(stripped):
        depth = stripped.split()[0].count(".") + stripped.split()[0].count("-")
        return min(depth + 1, 6)
    if stripped.isupper() and len(words) >= 2:
        return 2
    # Title Case with no terminal punctuation, e.g. "Experimental Results".
    alpha_words = [word for word in words if word[:1].isalpha()]
    if len(alpha_words) >= 2 and all(word[:1].isupper() for word in alpha_words):
        return 3
    return None


def _blocks_from_plain_text(text: str, page: int | None = None) -> list[Block]:
    """Split plain text into blocks, promoting lines that look like headings."""
    blocks: list[Block] = []
    for paragraph in _BLANK_LINE_RE.split(text):
        cleaned = paragraph.strip()
        if not cleaned:
            continue
        lines = [line for line in cleaned.splitlines() if line.strip()]
        # A lone line in its own paragraph is the only heading candidate; a
        # short line inside a wrapped paragraph is just a short line.
        if len(lines) == 1:
            level = _looks_like_heading(lines[0])
            if level is not None:
                blocks.append(
                    Block(text=lines[0].strip(), kind="heading", level=level, page=page)
                )
                continue
        blocks.append(Block(text=" ".join(line.strip() for line in lines), page=page))
    return blocks


# ------------------------------------------------------------------- pdf


def load_pdf(path: Path) -> NormalizedDoc:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise LoaderError("pypdf is required to read PDFs (uv sync --extra rag)") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise LoaderError(f"Could not open PDF {path.name}: {exc}") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise LoaderError(f"PDF {path.name} is encrypted") from exc

    blocks: list[Block] = []
    empty_pages = 0
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            logger.debug("Page %d of %s failed to extract: %s", index, path.name, exc)
            text = ""
        text = _repair_pdf_linebreaks(text)
        if not text.strip():
            empty_pages += 1
            continue
        blocks.extend(_blocks_from_plain_text(text, page=index))

    page_count = len(reader.pages)
    metadata: dict[str, object] = {"page_count": page_count, "empty_pages": empty_pages}
    title = None
    try:
        info = reader.metadata or {}
        title = (info.get("/Title") or "").strip() or None
        for key, name in (("/Author", "author"), ("/Producer", "producer")):
            value = (info.get(key) or "").strip()
            if value:
                metadata[name] = value
    except Exception:
        logger.debug("Could not read PDF metadata for %s", path.name)

    # A PDF that yields almost no text is a scan. Say so plainly rather than
    # silently ingesting an empty document that will never be retrievable.
    if page_count and empty_pages / page_count > 0.8:
        raise LoaderError(
            f"{path.name} has no extractable text on {empty_pages}/{page_count} pages. "
            "It is probably a scan; OCR it first (e.g. `ocrmypdf in.pdf out.pdf`)."
        )

    return NormalizedDoc(
        source_path=str(path),
        doctype="pdf",
        blocks=blocks,
        title=title,
        page_count=page_count,
        metadata=metadata,
    )


def _repair_pdf_linebreaks(text: str) -> str:
    """Undo the hard wrapping PDF extraction produces.

    Column layout gives every line its own newline, which would make each
    line a separate block. Hyphens split across lines are rejoined; other
    single newlines become spaces, while blank lines stay as paragraph breaks.
    """
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    return re.sub(r"[ \t]{2,}", " ", text)


# ------------------------------------------------------------------ docx


def load_docx(path: Path) -> NormalizedDoc:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise LoaderError("python-docx is required to read .docx (uv sync --extra rag)") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise LoaderError(f"Could not open {path.name}: {exc}") from exc

    blocks: list[Block] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        level = _docx_heading_level(style)
        if level is not None:
            blocks.append(Block(text=text, kind="heading", level=level))
        elif style.startswith("List"):
            blocks.append(Block(text=text, kind="list"))
        else:
            blocks.append(Block(text=text))

    for table in document.tables:
        rendered = _render_table([[cell.text.strip() for cell in row.cells] for row in table.rows])
        if rendered:
            blocks.append(Block(text=rendered, kind="table"))

    core = document.core_properties
    metadata: dict[str, object] = {}
    for attribute in ("author", "subject", "category", "comments"):
        value = getattr(core, attribute, None)
        if value:
            metadata[attribute] = str(value)
    if core.created:
        metadata["created"] = core.created.isoformat()

    return NormalizedDoc(
        source_path=str(path),
        doctype="docx",
        blocks=blocks,
        title=(core.title or "").strip() or None,
        metadata=metadata,
    )


def _docx_heading_level(style_name: str) -> int | None:
    if style_name == "Title":
        return 1
    match = re.match(r"^Heading (\d+)$", style_name)
    if match:
        return min(int(match.group(1)), 6)
    return None


def _render_table(rows: list[list[str]]) -> str:
    """Render a table as pipe-delimited Markdown.

    Tables are kept as text rather than dropped: a specification's limits and
    a report's numbers usually live in one, and losing them loses the answer.
    """
    cleaned = [[cell.replace("\n", " ").strip() for cell in row] for row in rows if any(row)]
    if not cleaned:
        return ""
    header, *body = cleaned
    lines = [" | ".join(header), " | ".join("---" for _ in header)]
    lines.extend(" | ".join(row) for row in body)
    return "\n".join(lines)


# ------------------------------------------------------------------ pptx


def load_pptx(path: Path) -> NormalizedDoc:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise LoaderError("python-pptx is required to read .pptx (uv sync --extra rag)") from exc

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise LoaderError(f"Could not open {path.name}: {exc}") from exc

    blocks: list[Block] = []
    slide_count = 0
    for number, slide in enumerate(presentation.slides, start=1):
        slide_count = number
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None

        title_text = (title_shape.text or "").strip() if title_shape is not None else ""
        blocks.append(
            Block(
                text=title_text or f"Slide {number}",
                kind="heading",
                level=2,
                page=number,
            )
        )

        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if getattr(shape, "has_table", False):
                table = shape.table
                rendered = _render_table(
                    [[cell.text.strip() for cell in row.cells] for row in table.rows]
                )
                if rendered:
                    blocks.append(Block(text=rendered, kind="table", page=number))
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    blocks.append(Block(text=text, kind="list", page=number))

        notes = _slide_notes(slide)
        if notes:
            blocks.append(Block(text=f"Speaker notes: {notes}", kind="caption", page=number))

    return NormalizedDoc(
        source_path=str(path),
        doctype="pptx",
        blocks=blocks,
        title=(presentation.core_properties.title or "").strip() or None,
        page_count=slide_count,
        metadata={"slide_count": slide_count},
    )


def _slide_notes(slide: object) -> str:
    try:
        if not slide.has_notes_slide:  # type: ignore[attr-defined]
            return ""
        return (slide.notes_slide.notes_text_frame.text or "").strip()  # type: ignore[attr-defined]
    except Exception:
        return ""


# ------------------------------------------------------------------ html


def _html_title(html_text: str) -> str | None:
    try:
        tree = lxml_html.fromstring(html_text)
    except Exception:
        return None
    nodes = tree.xpath("//title/text()")
    if not nodes:
        return None
    return str(nodes[0]).strip() or None


def blocks_from_html(html_text: str) -> tuple[list[Block], str | None]:
    """Extract structured blocks and a title from an HTML fragment."""
    from mcp_fetch_server.converters import sanitize_html

    # The sanitiser strips <head> along with scripts and styles, so the title
    # has to be read from the original markup before cleaning.
    title = _html_title(html_text)

    cleaned = sanitize_html(html_text) or html_text
    try:
        tree = lxml_html.fromstring(cleaned)
    except Exception as exc:
        raise LoaderError(f"Could not parse HTML: {exc}") from exc

    blocks: list[Block] = []
    body = tree.find("body")
    root = body if body is not None else tree

    selector = ".//h1 | .//h2 | .//h3 | .//h4 | .//h5 | .//h6 | .//p | .//li | .//pre | .//table"
    for node in root.xpath(selector):
        tag = str(node.tag).lower()
        if tag.startswith("h") and len(tag) == 2 and tag[1].isdigit():
            text = node.text_content().strip()
            if text:
                blocks.append(Block(text=text, kind="heading", level=int(tag[1])))
        elif tag == "pre":
            text = node.text_content().rstrip()
            if text.strip():
                blocks.append(Block(text=text, kind="code"))
        elif tag == "table":
            rows = [
                [cell.text_content().strip() for cell in row.xpath("./th | ./td")]
                for row in node.xpath(".//tr")
            ]
            rendered = _render_table(rows)
            if rendered:
                blocks.append(Block(text=rendered, kind="table"))
        elif tag == "li":
            text = " ".join(node.text_content().split())
            if text:
                blocks.append(Block(text=text, kind="list"))
        else:
            text = " ".join(node.text_content().split())
            if text:
                blocks.append(Block(text=text))

    if not blocks:
        text = " ".join(root.text_content().split())
        if text:
            blocks.append(Block(text=text))
    return blocks, title


def extract_html_links(html_text: str) -> list[tuple[str, str]]:
    """Collect (href, anchor text) pairs, so a mirrored site keeps its graph.

    An archived website is only browsable if its internal links survive
    ingestion; without them ``extract_links`` has nothing to walk and the
    corpus is a pile of disconnected pages.
    """
    try:
        tree = lxml_html.fromstring(html_text)
    except Exception:
        return []

    seen: set[str] = set()
    links: list[tuple[str, str]] = []
    for node in tree.xpath("//a[@href]"):
        href = str(node.get("href") or "").strip()
        if not href or href.startswith(("#", "javascript:", "mailto:", "data:")):
            continue
        if href in seen:
            continue
        seen.add(href)
        links.append((href, " ".join(node.text_content().split())[:200]))
    return links


def load_html(path: Path) -> NormalizedDoc:
    raw = _read_text(path)
    blocks, title = blocks_from_html(raw)
    return NormalizedDoc(
        source_path=str(path),
        doctype="html",
        blocks=blocks,
        title=title,
        metadata={"links": extract_html_links(raw)},
    )


# ------------------------------------------------------------------ epub


def load_epub(path: Path) -> NormalizedDoc:
    """Read an EPUB by unzipping it and walking the spine in reading order."""
    try:
        archive = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as exc:
        raise LoaderError(f"Could not open EPUB {path.name}: {exc}") from exc

    blocks: list[Block] = []
    title: str | None = None
    with archive:
        opf_path = _epub_opf_path(archive)
        opf_dir = posixpath.dirname(opf_path)
        try:
            opf = etree.fromstring(archive.read(opf_path))
        except (KeyError, etree.XMLSyntaxError) as exc:
            raise LoaderError(f"EPUB {path.name} has an unreadable package file: {exc}") from exc

        namespaces = {
            "opf": "http://www.idpf.org/2007/opf",
            "dc": "http://purl.org/dc/elements/1.1/",
        }
        title_nodes = opf.xpath("//dc:title/text()", namespaces=namespaces)
        if title_nodes:
            title = str(title_nodes[0]).strip() or None

        manifest = {
            str(item.get("id")): str(item.get("href"))
            for item in opf.xpath("//opf:manifest/opf:item", namespaces=namespaces)
        }
        spine_ids = [
            str(ref.get("idref"))
            for ref in opf.xpath("//opf:spine/opf:itemref", namespaces=namespaces)
        ]

        for order, item_id in enumerate(spine_ids, start=1):
            href = manifest.get(item_id)
            if not href:
                continue
            name = posixpath.normpath(posixpath.join(opf_dir, href)) if opf_dir else href
            try:
                content = archive.read(name)
            except KeyError:
                logger.debug("EPUB %s references a missing item: %s", path.name, name)
                continue
            try:
                section_blocks, _ = blocks_from_html(content.decode("utf-8", errors="replace"))
            except LoaderError:
                continue
            for block in section_blocks:
                block.page = order
            blocks.extend(section_blocks)

    if not blocks:
        raise LoaderError(f"EPUB {path.name} contained no readable text")

    return NormalizedDoc(
        source_path=str(path),
        doctype="epub",
        blocks=blocks,
        title=title,
        page_count=max((block.page or 0) for block in blocks) or None,
        metadata={"sections": len({block.page for block in blocks})},
    )


def _epub_opf_path(archive: zipfile.ZipFile) -> str:
    try:
        container = etree.fromstring(archive.read("META-INF/container.xml"))
    except (KeyError, etree.XMLSyntaxError) as exc:
        raise LoaderError("EPUB is missing META-INF/container.xml") from exc
    namespaces = {"c": "urn:oasis:names:tc:opendocument:xmlns:container"}
    roots = container.xpath("//c:rootfile/@full-path", namespaces=namespaces)
    if not roots:
        raise LoaderError("EPUB container.xml declares no rootfile")
    return str(roots[0])


# -------------------------------------------------------------- markdown


def load_markdown(path: Path) -> NormalizedDoc:
    text = _read_text(path)
    blocks: list[Block] = []
    in_fence = False
    fence_lines: list[str] = []
    paragraph: list[str] = []

    def flush_paragraph() -> None:
        if paragraph:
            blocks.append(Block(text=" ".join(paragraph).strip()))
            paragraph.clear()

    for line in text.splitlines():
        if _MD_FENCE_RE.match(line):
            if in_fence:
                blocks.append(Block(text="\n".join(fence_lines), kind="code"))
                fence_lines.clear()
                in_fence = False
            else:
                flush_paragraph()
                in_fence = True
            continue
        if in_fence:
            fence_lines.append(line)
            continue

        heading = _MD_HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            blocks.append(
                Block(
                    text=heading.group(2).strip(),
                    kind="heading",
                    level=len(heading.group(1)),
                )
            )
            continue

        if not line.strip():
            flush_paragraph()
            continue

        stripped = line.strip()
        if stripped.startswith(("- ", "* ", "+ ")) or re.match(r"^\d+[.)]\s", stripped):
            flush_paragraph()
            blocks.append(Block(text=stripped, kind="list"))
            continue

        paragraph.append(stripped)

    if in_fence and fence_lines:
        blocks.append(Block(text="\n".join(fence_lines), kind="code"))
    flush_paragraph()

    return NormalizedDoc(source_path=str(path), doctype="markdown", blocks=blocks)


def load_text(path: Path) -> NormalizedDoc:
    return NormalizedDoc(
        source_path=str(path),
        doctype="text",
        blocks=_blocks_from_plain_text(_read_text(path)),
    )


def _read_text(path: Path) -> str:
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise LoaderError(f"Could not read {path.name}: {exc}") from exc
    for encoding in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


# ---------------------------------------------------------------- registry


LOADERS: dict[str, Callable[[Path], NormalizedDoc]] = {
    "pdf": load_pdf,
    "docx": load_docx,
    "pptx": load_pptx,
    "epub": load_epub,
    "html": load_html,
    "markdown": load_markdown,
    "text": load_text,
}


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTENSIONS


def load_document(path: Path) -> NormalizedDoc:
    """Load any supported file into a :class:`NormalizedDoc`."""
    if not path.is_file():
        raise LoaderError(f"Not a file: {path}")

    doctype = EXTENSION_DOCTYPES.get(path.suffix.lower())
    if doctype is None:
        raise UnsupportedFormatError(
            f"No loader for '{path.suffix}'. Supported: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    document = LOADERS[doctype](path)
    if not document.blocks:
        raise LoaderError(f"{path.name} contained no readable text")

    document.language = detect_language(document.to_markdown())
    document.metadata.setdefault("byte_size", path.stat().st_size)
    return document
